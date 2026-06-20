"""
LNG BOR 예측 엑셀 모델 검증 스크립트
--------------------------------------
generate_bor_model.py로 생성된 LNG_BOR_Model.xlsx의 구조, 수식, 수치,
드롭다운, XML 수식 보존을 검사합니다.

실행 방법:
    python validate_model.py

검증 통과 시 "✅ 모든 검증 통과" 출력 후 exit code 0.
실패 시 오류 메시지 출력 후 exit code 1.
"""

import hashlib
import math
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook

from generate_bor_model import INPUT_ROWS, main as generate_workbook

MODEL_PATH = Path("LNG_BOR_Model.xlsx")

# 워크시트 탭 이름 (영문 ASCII)
REQUIRED_SHEETS = ["Input", "PropertyDB", "BOR_Calc", "Measure"]

# 기본 입력값(직경 4 m, 길이 10 m, 충전율 90%, Perlite(진공), 외기 25 ℃,
# Methane 90 / Ethane 7 / Propane 3)으로 계산한 BOR는 약 0.03 %/day여야 한다.
MIN_EXPECTED_BOR = 0.01
MAX_EXPECTED_BOR = 0.10
INSULATION_ROWS = range(4, 10)
COMPONENT_ROWS = range(13, 19)
PIPE_FACTOR_ROWS = range(22, 25)
HIGH_VACUUM_THRESHOLD_PA = 10
MEDIUM_VACUUM_THRESHOLD_PA = 100
HIGH_VACUUM_FACTOR = 0.7
MEDIUM_VACUUM_FACTOR = 0.85
ATMOSPHERIC_VACUUM_FACTOR = 1.0
REFERENCE_PRESSURE_BAR = 1.013
PRESSURE_CORRECTION_SLOPE = 0.05
MIN_PRESSURE_CORRECTION = 0.8

# OOXML 네임스페이스
NS_SPREADSHEET = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _norm_formula(value):
    """수식 문자열을 공백 제거 + 대문자 변환해 비교 가능하게 정규화."""
    if not isinstance(value, str):
        return value
    return re.sub(r"\s+", "", value).upper()


def _find_input_composition_rows(input_ws):
    """Input 시트에서 LNG 조성 3개 행(성분명/몰분율)을 찾아 행 번호 리스트로 반환."""
    section_row = None
    for row in range(1, input_ws.max_row + 1):
        value = input_ws.cell(row, 1).value
        if isinstance(value, str) and "▶ LNG 조성" in value:
            section_row = row
            break
    if section_row is None:
        return None

    rows = []
    for row in range(section_row + 1, input_ws.max_row + 1):
        name = input_ws.cell(row, 1).value
        frac = input_ws.cell(row, 2).value
        if isinstance(name, str) and isinstance(frac, (int, float)):
            rows.append(row)
            if len(rows) == 3:
                break
    return rows if len(rows) == 3 else None


def file_sha256(path):
    """파일 내용을 SHA-256으로 계산."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def validate_sheets(wb):
    """필수 시트 존재 여부 확인."""
    missing = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
    if missing:
        print(f"❌ 누락된 시트: {missing}")
        return False
    print(f"  ✔ 시트 구성 확인: {REQUIRED_SHEETS}")
    return True


def validate_formula_references(wb):
    """BOR_Calc LNG 조성 계산 수식이 올바른 열을 참조하는지 확인."""
    input_ws = wb["Input"]
    ws = wb["BOR_Calc"]
    errors = []
    comp_rows = _find_input_composition_rows(input_ws)
    if not comp_rows:
        print("❌ Input 시트에서 LNG 조성(성분 이름/몰분율) 행 3개를 찾지 못했습니다.")
        return False

    lookup_cells = {
        "B16": (3, comp_rows[0]),
        "B17": (3, comp_rows[1]),
        "B18": (3, comp_rows[2]),
        "B20": (4, comp_rows[0]),
        "B21": (4, comp_rows[1]),
        "B22": (4, comp_rows[2]),
        "B24": (5, comp_rows[0]),
        "B25": (5, comp_rows[1]),
        "B26": (5, comp_rows[2]),
    }
    for cell_ref, (db_col, input_row) in lookup_cells.items():
        expected = f"=VLOOKUP(Input!A{input_row},PropertyDB!A13:E18,{db_col},FALSE)"
        actual = ws[cell_ref].value
        if _norm_formula(actual) != _norm_formula(expected):
            errors.append(f"{cell_ref} lookup 수식 오류: {actual!r} != {expected!r}")

    r1, r2, r3 = comp_rows
    weighted_expected = {
        "B19": f"=(B16*Input!B{r1}+B17*Input!B{r2}+B18*Input!B{r3})/(Input!B{r1}+Input!B{r2}+Input!B{r3})",
        "B23": f"=(B20*Input!B{r1}+B21*Input!B{r2}+B22*Input!B{r3})/(Input!B{r1}+Input!B{r2}+Input!B{r3})",
        "B27": f"=(B24*Input!B{r1}+B25*Input!B{r2}+B26*Input!B{r3})/(Input!B{r1}+Input!B{r2}+Input!B{r3})",
    }
    for cell_ref, expected in weighted_expected.items():
        actual = ws[cell_ref].value
        if _norm_formula(actual) != _norm_formula(expected):
            errors.append(f"{cell_ref} 가중평균 수식 오류: {actual!r} != {expected!r}")

    if errors:
        for e in errors:
            print(f"❌ {e}")
        return False

    print("  ✔ BOR_Calc LNG 조성 수식 열 참조 확인 (성분명=A열, 몰분율=B열)")
    return True


def validate_formulas_xml(path):
    """xlsx ZIP 내부 sheet XML들에 <f> 태그가 존재하는지 확인."""
    total_formula_tags = 0
    worksheet_files = []

    with zipfile.ZipFile(path, "r") as zf:
        for name in zf.namelist():
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                worksheet_files.append(name)
                sheet_root = ET.fromstring(zf.read(name))
                total_formula_tags += len(
                    list(sheet_root.iter(f"{{{NS_SPREADSHEET}}}f"))
                )

    if not worksheet_files:
        print("❌ xlsx 내부에서 xl/worksheets/sheet*.xml 파일을 찾을 수 없습니다.")
        return False
    if total_formula_tags == 0:
        print("❌ xlsx 내부 시트 XML들에 <f> 수식 태그가 없습니다.")
        return False

    print(f"  ✔ XML 수식 태그 확인: {len(worksheet_files)}개 시트 XML, <f> 태그 {total_formula_tags}개")
    return True


def validate_dropdown(wb):
    """Input 시트에 드롭다운(데이터 유효성) 검사가 존재하는지 확인."""
    ws = wb["Input"]
    data_validation_rules = list(ws.data_validations.dataValidation)
    if not data_validation_rules:
        print("❌ Input 시트에 드롭다운(데이터 유효성) 검사가 없습니다.")
        return False
    ranges = {str(rule.sqref) for rule in data_validation_rules}
    expected_ranges = {
        f"B{INPUT_ROWS['insulation']}",
        f"B{INPUT_ROWS['pipe_state']}",
        f"B{INPUT_ROWS['comp1']}:B{INPUT_ROWS['comp3']}",
    }
    if not expected_ranges.issubset(ranges):
        print(f"❌ Input 드롭다운 범위가 예상과 다릅니다. 현재: {sorted(ranges)}")
        return False
    print(f"  ✔ Input 드롭다운 유효성 검사 확인: {len(data_validation_rules)}개 존재")
    return True


def validate_bor_range(wb):
    """기본 입력값으로 독립 계산한 BOR 값이 예상 범위인지 확인."""
    input_ws = wb["Input"]
    prop_ws = wb["PropertyDB"]

    insulation = {
        prop_ws.cell(row, 1).value: float(prop_ws.cell(row, 2).value)
        for row in INSULATION_ROWS
    }
    components = {
        prop_ws.cell(row, 1).value: {
            "hvap": float(prop_ws.cell(row, 3).value),
            "rho": float(prop_ws.cell(row, 4).value),
            "tbp": float(prop_ws.cell(row, 5).value),
        }
        for row in COMPONENT_ROWS
    }
    pipe_factors = {
        prop_ws.cell(row, 1).value: float(prop_ws.cell(row, 2).value)
        for row in PIPE_FACTOR_ROWS
    }

    diameter = float(input_ws[f"B{INPUT_ROWS['diameter']}"].value)
    length = float(input_ws[f"B{INPUT_ROWS['length']}"].value)
    fill_ratio = float(input_ws[f"B{INPUT_ROWS['fill_ratio']}"].value)
    pressure = float(input_ws[f"B{INPUT_ROWS['pressure']}"].value)
    ambient_temp = float(input_ws[f"B{INPUT_ROWS['ambient_temp']}"].value)
    insulation_name = input_ws[f"B{INPUT_ROWS['insulation']}"].value
    insulation_thickness = float(input_ws[f"B{INPUT_ROWS['insulation_thickness']}"].value)
    vacuum_pa = float(input_ws[f"B{INPUT_ROWS['vacuum_pa']}"].value)
    pipe_state = input_ws[f"B{INPUT_ROWS['pipe_state']}"].value
    pipe_heat_ratio = float(input_ws[f"B{INPUT_ROWS['pipe_ratio']}"].value)

    comp_rows = _find_input_composition_rows(input_ws)
    if not comp_rows:
        print("❌ Input 시트에서 LNG 조성 행을 찾지 못해 BOR 독립 계산을 수행할 수 없습니다.")
        return False
    composition = [
        (input_ws[f"A{comp_rows[0]}"].value, float(input_ws[f"B{comp_rows[0]}"].value)),
        (input_ws[f"A{comp_rows[1]}"].value, float(input_ws[f"B{comp_rows[1]}"].value)),
        (input_ws[f"A{comp_rows[2]}"].value, float(input_ws[f"B{comp_rows[2]}"].value)),
    ]
    total_fraction = sum(fraction for _, fraction in composition)

    avg_hvap = sum(components[name]["hvap"] * fraction for name, fraction in composition) / total_fraction
    avg_rho = sum(components[name]["rho"] * fraction for name, fraction in composition) / total_fraction
    avg_tbp = sum(components[name]["tbp"] * fraction for name, fraction in composition) / total_fraction

    radius = diameter / 2
    area_total = (2 * math.pi * radius * length) + (4 * math.pi * radius**2)
    volume_tank = (math.pi * radius**2 * length) + ((4 / 3) * math.pi * radius**3)
    volume_lng = volume_tank * fill_ratio / 100

    vacuum_factor = (
        HIGH_VACUUM_FACTOR
        if vacuum_pa < HIGH_VACUUM_THRESHOLD_PA
        else MEDIUM_VACUUM_FACTOR
        if vacuum_pa < MEDIUM_VACUUM_THRESHOLD_PA
        else ATMOSPHERIC_VACUUM_FACTOR
    )
    k_eff = insulation[insulation_name] * vacuum_factor
    thermal_resistance = insulation_thickness / (k_eff * area_total)

    pressure_correction = max(
        MIN_PRESSURE_CORRECTION,
        1 - ((pressure - REFERENCE_PRESSURE_BAR) * PRESSURE_CORRECTION_SLOPE),
    )
    delta_t = ambient_temp - avg_tbp
    q_tank = delta_t / thermal_resistance
    q_pipe = q_tank * (pipe_heat_ratio / 100) * pipe_factors[pipe_state]
    q_total = q_tank + q_pipe

    bog_kg_s = q_total / (avg_hvap * 1000)
    bog_kg_day = bog_kg_s * 86400 * pressure_correction
    lng_mass = volume_lng * avg_rho
    bor_percent_day = bog_kg_day / lng_mass * 100

    if not (MIN_EXPECTED_BOR <= bor_percent_day <= MAX_EXPECTED_BOR):
        print(f"❌ 독립 계산 BOR 값이 예상 범위를 벗어났습니다: {bor_percent_day:.5f} %/day")
        return False

    print(f"  ✔ 독립 계산 BOR 확인: {bor_percent_day:.5f} %/day")
    return True


def validate_note_cells_no_formula_prefix(wb):
    """BOR_Calc/Measure D열 note 텍스트가 '='로 시작하지 않는지 확인."""
    errors = []
    for sheet_name in ["BOR_Calc", "Measure"]:
        ws = wb[sheet_name]
        for row in range(3, ws.max_row + 1):
            value = ws.cell(row, 4).value
            if isinstance(value, str) and value.startswith("="):
                errors.append(f"{sheet_name}!D{row} note가 '='로 시작: {value!r}")

    if errors:
        for e in errors:
            print(f"❌ {e}")
        return False

    print("  ✔ BOR_Calc/Measure D열 note '=' 시작 문자열 없음")
    return True


def validate_input_volume_rows(wb):
    """Input 시트에 체적 자동 계산 행이 존재하고 B열이 수식인지 확인."""
    ws = wb["Input"]
    row_tank = row_lng = None
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if label == "탱크 내부 체적 [m³]":
            row_tank = row
        elif label == "LNG 체적 [m³]":
            row_lng = row

    if row_tank is None or row_lng is None:
        print(f"❌ Input 체적 행 누락: 탱크={row_tank}, LNG={row_lng}")
        return False

    tank_formula = ws.cell(row_tank, 2).value
    lng_formula = ws.cell(row_lng, 2).value
    if not (isinstance(tank_formula, str) and tank_formula.startswith("=")):
        print(f"❌ 탱크 내부 체적 B{row_tank}가 수식이 아닙니다: {tank_formula!r}")
        return False
    if not (isinstance(lng_formula, str) and lng_formula.startswith("=")):
        print(f"❌ LNG 체적 B{row_lng}가 수식이 아닙니다: {lng_formula!r}")
        return False

    print(f"  ✔ Input 체적 자동 계산 행 확인: A{row_tank}, A{row_lng}")
    return True


def main():
    print(f"🔍 검증 시작: {MODEL_PATH}\n")

    before_hash = file_sha256(MODEL_PATH) if MODEL_PATH.exists() else None

    print("[0] 엑셀 모델 재생성")
    generate_workbook()
    if not MODEL_PATH.exists():
        print(f"❌ generate_bor_model.main() 실행 후 파일이 생성되지 않았습니다: {MODEL_PATH}")
        sys.exit(1)
    if before_hash is not None:
        print(f"  ✔ 파일 재생성 확인: {MODEL_PATH} (SHA-256 {before_hash[:8]} → {file_sha256(MODEL_PATH)[:8]})")
    else:
        print(f"  ✔ 파일 생성 확인: {MODEL_PATH}")

    try:
        wb = load_workbook(MODEL_PATH)
    except FileNotFoundError:
        print(f"❌ 파일을 찾을 수 없습니다: {MODEL_PATH}")
        print("   먼저 python generate_bor_model.py 를 실행하세요.")
        sys.exit(1)

    results = []

    print("[1] 시트 구성 확인")
    results.append(validate_sheets(wb))

    print("[2] BOR_Calc 수식 열 참조 확인")
    results.append(validate_formula_references(wb))

    print("[3] BOR_Calc/Measure D열 note 회귀 검사")
    results.append(validate_note_cells_no_formula_prefix(wb))

    print("[4] Input 체적 자동계산 행 확인")
    results.append(validate_input_volume_rows(wb))

    print("[5] 시트 XML 수식 태그 확인")
    results.append(validate_formulas_xml(MODEL_PATH))

    print("[6] 기본 입력값 독립 계산 BOR 확인")
    results.append(validate_bor_range(wb))

    print("[7] Input 시트 드롭다운 유효성 확인")
    results.append(validate_dropdown(wb))

    print()
    if all(results):
        print("✅ 모든 검증 통과")
        sys.exit(0)
    else:
        failed = sum(1 for r in results if not r)
        print(f"❌ 검증 실패: {failed}개 항목 오류")
        sys.exit(1)


if __name__ == "__main__":
    main()
